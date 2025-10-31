"""
AI 콘텐츠 생성 모듈
Google Gemini API를 활용한 블로그 포스트 자동 생성
"""

import os
import time
from dotenv import load_dotenv
import google.generativeai as genai
from openpyxl import load_workbook

# 환경 변수 로드
load_dotenv()


class AIContentGenerator:
    def __init__(self):
        """AI 콘텐츠 생성기 초기화"""
        self.api_key = os.getenv('GEMINI_API_KEY')

        if not self.api_key:
            raise ValueError("환경 변수에 GEMINI_API_KEY를 설정해주세요.")

        # Gemini API 설정
        genai.configure(api_key=self.api_key)
        self.model = genai.GenerativeModel('gemini-pro')

    def generate_blog_post(self, keyword, tone="친근하고 정보적인", length="중간"):
        """
        키워드를 바탕으로 블로그 포스트 생성

        Args:
            keyword (str): 블로그 포스트 주제 키워드
            tone (str): 글의 톤 (예: "친근하고 정보적인", "전문적인", "유머러스한")
            length (str): 글의 길이 ("짧은", "중간", "긴")

        Returns:
            dict: {'title': 제목, 'content': 본문 내용 리스트}
        """
        print(f"\n키워드 '{keyword}'에 대한 블로그 포스트 생성 중...")

        # 길이에 따른 문단 수 설정
        paragraph_count = {
            "짧은": 3,
            "중간": 5,
            "긴": 8
        }.get(length, 5)

        # 프롬프트 생성
        prompt = f"""
당신은 전문 블로그 작가입니다. 다음 키워드에 대한 네이버 블로그 포스트를 작성해주세요.

키워드: {keyword}
글의 톤: {tone}
문단 수: 약 {paragraph_count}개

다음 형식으로 작성해주세요:

[제목]
(매력적이고 클릭하고 싶은 제목을 한 줄로 작성)

[본문]
(각 문단은 2-3 문장으로 구성하며, 독자에게 유용한 정보를 제공하세요)
(문단마다 줄바꿈을 넣어주세요)

주의사항:
- 제목은 [제목] 태그 다음 줄에 작성
- 본문은 [본문] 태그 다음부터 작성
- SEO에 최적화된 내용으로 작성
- 독자의 흥미를 유발하는 내용 포함
- 자연스러운 한글 표현 사용
"""

        try:
            # API 호출
            response = self.model.generate_content(prompt)

            # 응답 파싱
            generated_text = response.text

            # 제목과 본문 분리
            title = ""
            content_lines = []

            lines = generated_text.strip().split('\n')

            in_title_section = False
            in_content_section = False

            for line in lines:
                line = line.strip()

                if '[제목]' in line or '제목:' in line:
                    in_title_section = True
                    in_content_section = False
                    continue
                elif '[본문]' in line or '본문:' in line:
                    in_title_section = False
                    in_content_section = True
                    continue

                if in_title_section and line and not title:
                    title = line
                elif in_content_section and line:
                    content_lines.append(line)

            # 제목이 없으면 기본 제목 설정
            if not title:
                title = f"[{keyword}] 에 대한 모든 것"

            # 본문이 비어있으면 기본 내용 설정
            if not content_lines:
                content_lines = [
                    f"{keyword}에 대해 알아보겠습니다.",
                    "다양한 정보를 제공해드립니다.",
                    "유용한 내용이 되었으면 좋겠습니다."
                ]

            print(f"✓ 제목: {title}")
            print(f"✓ 본문: {len(content_lines)}개 문단 생성 완료")

            return {
                'title': title,
                'content': content_lines
            }

        except Exception as e:
            print(f"AI 콘텐츠 생성 중 오류 발생: {str(e)}")
            # 오류 발생 시 기본 콘텐츠 반환
            return {
                'title': f"[{keyword}] 자동 생성 포스트",
                'content': [
                    f"{keyword}에 대한 정보를 제공합니다.",
                    "AI가 생성한 콘텐츠입니다.",
                    "내용을 확인하고 수정해주세요."
                ]
            }

    def generate_from_excel(self, excel_file_path, sheet_name='Sheet1', keyword_column='A'):
        """
        엑셀 파일에서 키워드를 읽어 대량으로 콘텐츠 생성

        Args:
            excel_file_path (str): 엑셀 파일 경로
            sheet_name (str): 시트 이름
            keyword_column (str): 키워드가 있는 열 (예: 'A', 'B', 'C')

        Returns:
            list: [{'keyword': 키워드, 'title': 제목, 'content': 본문}, ...]
        """
        print(f"\n엑셀 파일 '{excel_file_path}'에서 키워드 읽기 중...")

        try:
            # 엑셀 파일 로드
            workbook = load_workbook(excel_file_path)
            sheet = workbook[sheet_name]

            # 키워드 추출
            keywords = []
            for row in sheet.iter_rows(min_row=2, values_only=True):  # 헤더 제외
                keyword = row[ord(keyword_column.upper()) - ord('A')]
                if keyword:
                    keywords.append(str(keyword).strip())

            print(f"✓ {len(keywords)}개의 키워드를 찾았습니다.")

            # 각 키워드에 대해 콘텐츠 생성
            results = []
            for i, keyword in enumerate(keywords, 1):
                print(f"\n[{i}/{len(keywords)}] 진행 중...")

                post_data = self.generate_blog_post(keyword)
                results.append({
                    'keyword': keyword,
                    'title': post_data['title'],
                    'content': post_data['content']
                })

                # API 호출 제한을 고려한 대기 시간
                if i < len(keywords):
                    print("잠시 대기 중... (API 제한 방지)")
                    time.sleep(2)

            print(f"\n✓ 총 {len(results)}개의 블로그 포스트가 생성되었습니다.")
            return results

        except FileNotFoundError:
            print(f"오류: 엑셀 파일을 찾을 수 없습니다: {excel_file_path}")
            return []
        except Exception as e:
            print(f"엑셀 파일 처리 중 오류 발생: {str(e)}")
            return []


def main():
    """테스트용 메인 함수"""
    print("=" * 50)
    print("AI 콘텐츠 생성 모듈 테스트")
    print("=" * 50)

    try:
        generator = AIContentGenerator()

        # 사용자 입력 받기
        print("\n1. 단일 키워드 입력")
        print("2. 엑셀 파일에서 대량 생성")
        choice = input("선택하세요 (1 또는 2): ").strip()

        if choice == '1':
            keyword = input("\n키워드를 입력하세요: ").strip()
            if keyword:
                result = generator.generate_blog_post(keyword)
                print("\n" + "=" * 50)
                print("생성된 콘텐츠:")
                print("=" * 50)
                print(f"\n제목: {result['title']}")
                print(f"\n본문:")
                for line in result['content']:
                    print(f"  {line}")
            else:
                print("키워드를 입력해주세요.")

        elif choice == '2':
            excel_file = input("\n엑셀 파일 경로를 입력하세요: ").strip()
            if os.path.exists(excel_file):
                results = generator.generate_from_excel(excel_file)

                print("\n" + "=" * 50)
                print("생성된 콘텐츠 목록:")
                print("=" * 50)

                for i, result in enumerate(results, 1):
                    print(f"\n[{i}] 키워드: {result['keyword']}")
                    print(f"    제목: {result['title']}")
                    print(f"    본문: {len(result['content'])}개 문단")
            else:
                print("파일을 찾을 수 없습니다.")

        else:
            print("잘못된 선택입니다.")

    except ValueError as e:
        print(f"\n오류: {str(e)}")
        print("GEMINI_API_KEY를 .env 파일에 설정해주세요.")


if __name__ == "__main__":
    main()
