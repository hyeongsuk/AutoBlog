"""
엑셀 예제 파일 생성 스크립트
대량 블로그 포스팅용 키워드 목록 예제
"""

from openpyxl import Workbook


def create_example_excel():
    """예제 엑셀 파일 생성"""

    # 새 워크북 생성
    wb = Workbook()
    ws = wb.active
    ws.title = "키워드 목록"

    # 헤더 작성
    ws['A1'] = "키워드"

    # 예제 키워드 데이터
    example_keywords = [
        "파이썬 프로그래밍 시작하기",
        "웹 개발 트렌드 2024",
        "인공지능 활용 사례",
        "데이터 분석 기초",
        "블로그 운영 노하우"
    ]

    # 데이터 입력
    for i, keyword in enumerate(example_keywords, start=2):
        ws[f'A{i}'] = keyword

    # 파일 저장
    filename = "keywords_example.xlsx"
    wb.save(filename)

    print(f"예제 엑셀 파일이 생성되었습니다: {filename}")
    print(f"총 {len(example_keywords)}개의 키워드가 포함되어 있습니다.")
    print("\n포함된 키워드:")
    for i, keyword in enumerate(example_keywords, 1):
        print(f"  {i}. {keyword}")


if __name__ == "__main__":
    create_example_excel()
